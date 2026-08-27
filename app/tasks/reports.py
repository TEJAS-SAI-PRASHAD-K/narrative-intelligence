"""Report rendering.

CSV always works; PDF and PPTX need the `reports` extra. A missing renderer
fails the job with a message naming the extra to install rather than a stack
trace about an import.
"""

from __future__ import annotations

import csv
import io
import logging
from pathlib import Path
from typing import Any

from app.tasks.base import TrackedTask, report_progress
from app.tasks.celery_app import celery

log = logging.getLogger(__name__)


@celery.task(name="reports.render", base=TrackedTask, bind=True)
def render(self, *, job_id: str | None = None, project_id: str = "", **params: Any) -> dict:
    """Render a report to disk and record its path."""
    import uuid

    from sqlalchemy import text

    from app.config import get_api_settings
    from app.db import sync_session, utcnow
    from app.etl.service import resolve_project

    settings = get_api_settings()
    template = params.get("template", "exec_summary")
    fmt = params.get("format", "csv")
    report_id = uuid.uuid4()

    with sync_session() as session:
        resolved = resolve_project(session, project_id)
        session.execute(
            text(
                """
                INSERT INTO reports (id, project_id, job_id, template, format, status, params)
                VALUES (:id, :p, :job, :template, :format, 'running', cast(:params AS jsonb))
                """
            ),
            {
                "id": report_id,
                "p": resolved,
                "job": uuid.UUID(job_id) if job_id else None,
                "template": template,
                "format": fmt,
                "params": _dumps(params),
            },
        )
        report_progress(job_id, 0.2)

        rows = _gather(session, resolved, params)
        report_progress(job_id, 0.6)

        directory = Path(settings.reports_dir)
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / f"{report_id}.{fmt}"

        try:
            if fmt == "csv":
                _render_csv(path, rows)
            elif fmt == "pdf":
                _render_pdf(path, template, rows)
            elif fmt == "pptx":
                _render_pptx(path, template, rows)
            else:
                raise ValueError(f"unsupported format {fmt!r}")
        except Exception as exc:
            session.execute(
                text(
                    """
                    UPDATE reports SET status = 'failed', error = :error, finished_at = :now
                    WHERE id = :id
                    """
                ),
                {"id": report_id, "error": f"{type(exc).__name__}: {exc}"[:2000], "now": utcnow()},
            )
            raise

        session.execute(
            text(
                """
                UPDATE reports SET
                    status = 'succeeded', file_path = :path, size_bytes = :size,
                    finished_at = :now
                WHERE id = :id
                """
            ),
            {
                "id": report_id,
                "path": str(path),
                "size": path.stat().st_size,
                "now": utcnow(),
            },
        )

    result = {
        "report_id": str(report_id),
        "format": fmt,
        "rows": len(rows),
        "size_bytes": path.stat().st_size,
        "download_url": f"/api/v1/reports/{report_id}/download",
    }
    report_progress(job_id, 1.0, result=result)
    return result


def _gather(session, project_id, params: dict) -> list[dict[str, Any]]:
    """The narrative table every template is built from."""
    from sqlalchemy import text

    narrative_ids = params.get("narrative_ids") or []
    scope = "AND n.id = ANY(:ids)" if narrative_ids else ""
    import uuid as _uuid

    rows = session.execute(
        text(
            f"""
            SELECT n.display_id, n.title, n.claim, n.post_count, n.author_count,
                   n.engagement_total, n.date_start, n.date_end,
                   sc.priority, sc.fusion_score, sc.bot_like, sc.toxicity,
                   sc.negative_sentiment, sc.compass_risk, sc.scoring_version,
                   c.verification_status
            FROM narratives n
            LEFT JOIN narrative_scorecards sc ON sc.narrative_id = n.id
            LEFT JOIN compass_contexts c
                   ON c.narrative_id = n.id AND c.superseded_by IS NULL
            WHERE n.project_id = :p {scope}
            ORDER BY sc.fusion_score DESC NULLS LAST
            """  # noqa: S608 - scope is a fixed string, values are bound
        ),
        {
            "p": project_id,
            **({"ids": [_uuid.UUID(n) for n in narrative_ids]} if narrative_ids else {}),
        },
    ).all()
    return [dict(row._mapping) for row in rows]


def _render_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    """Raw data export.

    An empty result still writes a file with headers. A zero-byte download
    reads as a broken export; a header row with no data reads as "no narratives
    matched", which is the truth.
    """
    columns = [
        "display_id",
        "title",
        "claim",
        "priority",
        "fusion_score",
        "bot_like",
        "toxicity",
        "negative_sentiment",
        "compass_risk",
        "verification_status",
        "post_count",
        "author_count",
        "engagement_total",
        "date_start",
        "date_end",
        "scoring_version",
    ]
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({k: row.get(k) for k in columns})
    path.write_text(buffer.getvalue(), encoding="utf-8")


def _render_pdf(path: Path, template: str, rows: list[dict[str, Any]]) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise RuntimeError(
            "PDF rendering needs the 'reports' extra: pip install -e '.[reports]'"
        ) from exc

    styles = getSampleStyleSheet()
    document = SimpleDocTemplate(str(path), pagesize=landscape(A4))
    story = [
        Paragraph("Narrative Intelligence — narrative summary", styles["Title"]),
        Paragraph(
            f"{len(rows)} narrative(s), ordered by fusion score. Scores are "
            "percentile-ranked within this project; see docs/scoring.md.",
            styles["Normal"],
        ),
        Spacer(1, 16),
    ]

    limit = 40 if template == "exec_summary" else len(rows)
    data = [["ID", "Title", "Priority", "Fusion", "Posts", "Authors", "Compass"]]
    for row in rows[:limit]:
        data.append(
            [
                str(row.get("display_id") or ""),
                (row.get("title") or "")[:70],
                row.get("priority") or "",
                # An unscored narrative prints as an em dash, not 0.0. A zero
                # in a PDF a stakeholder reads is a claim that it was measured.
                ("—" if row.get("fusion_score") is None else f"{row['fusion_score']:.1f}"),
                str(row.get("post_count") or 0),
                str(row.get("author_count") or 0),
                row.get("verification_status") or "—",
            ]
        )

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story.append(table)
    document.build(story)


def _render_pptx(path: Path, template: str, rows: list[dict[str, Any]]) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Pt
    except ImportError as exc:
        raise RuntimeError(
            "PPTX rendering needs the 'reports' extra: pip install -e '.[reports]'"
        ) from exc

    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = "Narrative Intelligence"
    title_slide.placeholders[1].text = f"{len(rows)} narratives, ordered by fusion score"

    for row in rows[: (10 if template == "exec_summary" else 40)]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = (row.get("title") or "Untitled narrative")[:80]
        frame = slide.placeholders[1].text_frame
        score = row.get("fusion_score")
        frame.text = (
            f"Fusion score: {'not scored' if score is None else f'{score:.1f}'} "
            f"({row.get('priority') or 'unknown'} priority)"
        )
        for line in (
            f"Posts: {row.get('post_count') or 0}   Authors: {row.get('author_count') or 0}",
            f"Compass: {row.get('verification_status') or 'not checked'}",
            f"Claim: {(row.get('claim') or '—')[:160]}",
        ):
            paragraph = frame.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(14)

    presentation.save(str(path))


def _dumps(value: Any) -> str:
    import json

    return json.dumps(value, default=str)
